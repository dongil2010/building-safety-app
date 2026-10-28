"""SafetyMan.exe 디스어셈블로 확인한 가중치로 xlsx 층·상태·종합 점수를 재현한다.

exe 위치(이미지 베이스 0x400000):
  - 0x4664e0  부재 밀도 (구조형식 [obj+0x14] × 재료 [obj+0x18])
              라멘 RC: 기둥 0.9 · 내력벽 0.9 · 큰보 0.7 · 작은보 0.5 · 슬래브 0.3
              벽식: 내력벽 0.9 · 슬래브 0.5 / 무량판: 기둥 0.9 · 슬래브 0.9 / 강구조는 내력벽 0
  - 0x40d354  층 위치 밀도: r = (지상층수 - 층 [+1 지상층]) / 전체층수
              r<0.2 → 0.1, <0.4 → 0.3, <0.6 → 0.5, <0.8 → 0.7, 그 외 0.9
  - 0x454c2e  종합 밀도: 안전성 0.9 · 상태 0.7 · 변위 0.7
  - 0x463040  항목→부재 퍼지 규칙표 / 0x465400 Choquet / 0x465830 λ 탐색(-1~5, 0.001 간격)

실행: py -3 scripts/sfm/verify_xlsx_scores.py  (C:\\Users\\dlawo\\세이프티맨\\*.xlsx 전부 비교)
"""
import openpyxl, re, os, glob

def lam_measure(ws, l):
    if abs(l) < 1e-12:
        return sum(ws)
    p = 1.0
    for w in ws:
        p *= 1 + l * w
    return (p - 1) / l

def sugeno_lambda(ws):
    best_l, best_e = -1, float("inf")
    for i in range(6001):          # exe: λ = -1 … 5, step 0.001
        l = -1 + i * 1e-3
        e = abs(lam_measure(ws, l) - 1)
        if e < best_e:
            best_e, best_l = e, l
    return best_l

def choquet(vals, ws):
    pairs = [(v, w) for v, w in zip(vals, ws) if v >= 1e-5 and w > 0]
    if not pairs:
        return 0.0
    l = sugeno_lambda([w for _, w in pairs])
    pairs.sort(key=lambda p: p[0])
    r = pairs[0][0]
    for i in range(1, len(pairs)):
        r += (pairs[i][0] - pairs[i - 1][0]) * lam_measure([w for _, w in pairs[i:]], l)
    return r

def floor_density(f, above, below):  # exe 0x40d354
    total = above + below
    r = ((above - f) if f < 0 else (above - f + 1)) / total
    return 0.1 if r < 0.2 else 0.3 if r < 0.4 else 0.5 if r < 0.6 else 0.7 if r < 0.8 else 0.9

RC_W = {"기둥": 0.9, "내력벽": 0.9, "큰보": 0.7, "작은보": 0.5, "슬래브": 0.3}  # exe 0x4664e0 (라멘·RC)
COLS = {"기둥": "T", "내력벽": "AA", "큰보": "AH", "작은보": "AO", "슬래브": "AV"}

def num(s):
    m = re.search(r"([0-9]+\.?[0-9]*)", str(s))
    return float(m.group(1)) if m else None

def run(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Sheet1"]
    above = int(num(ws["N24"].value)); below = int(num(ws["N25"].value))
    floors = []
    for r in range(35, 80, 3):
        lab = ws[f"B{r}"].value
        if not lab or "층" not in str(lab) or ws[f"N{r+1}"].value != "상태":
            continue
        fl = int(re.match(r"\s*(-?\d+)층", str(lab)).group(1))
        mem = {m: ws[f"{c}{r+1}"].value for m, c in COLS.items()}
        vals = [float(mem[m]) if isinstance(mem[m], (int, float)) else 0 for m in RC_W]
        got = choquet(vals, list(RC_W.values()))
        want = num(ws[f"BQ{r+1}"].value)
        disp = num(ws[f"BX{r}"].value)
        floors.append((fl, got, want, disp, str(lab).split("\n")[-1]))
    fin = str(ws["N" + str(max(r for r in range(40, 60) if ws[f"B{r}"].value == "최종결과"))].value)
    st_want = float(re.search(r"상태평가:\s*([0-9.]+)", fin).group(1))
    tot_want = float(re.search(r"종합평가:\s*([0-9.]+)", fin).group(1))
    st = choquet([g for _, g, _, _, _ in floors], [floor_density(f, above, below) for f, *_ in floors])
    disp = next((d for *_, d, _ in floors if d), None)
    tot = choquet([0, st, disp or 0], [0.9, 0.7, 0.7])   # exe 0x454c2e: 안전성 0.9, 상태 0.7, 변위 0.7
    print(f"## {os.path.basename(path)}  (지상 {above} / 지하 {below})")
    for fl, g, w, d, typ in floors:
        print(f"   {fl:>3}층 {typ:<8} 밀도 {floor_density(fl, above, below)}  계산 {g:.2f}  xlsx {w}  {'OK' if abs(round(g,2)-w) < 0.006 else 'XX'}")
    print(f"   상태 계산 {st:.2f} / xlsx {st_want}  {'OK' if abs(round(st,2)-st_want) < 0.006 else 'XX'}")
    print(f"   종합 계산 {tot:.2f} / xlsx {tot_want}  (변위 {disp})  {'OK' if abs(round(tot,2)-tot_want) < 0.006 else 'XX'}")

for p in sorted(glob.glob(r"C:\Users\dlawo\세이프티맨\*.xlsx")):
    try:
        run(p)
    except Exception as e:
        print("!!", os.path.basename(p), e)
