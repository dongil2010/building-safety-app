"""작성 시나리오 검증: 2026 원본 .sfm → 엑셀 → 탄산화 칸만 손으로 고치듯 수정 → .sfm
결과가 SafetyMan 이 직접 저장한 「2026종합평가 탄산화 수정본.sfm」과 바이트까지 같아야 한다."""
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
import sfm_format as F
import sfm_excel as X

DIR = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\dlawo\세이프티맨"
src = F.read(open(os.path.join(DIR, "2026종합평가.sfm"), "rb").read())
want_raw = open(os.path.join(DIR, "2026종합평가 탄산화 수정본.sfm"), "rb").read()
want = F.read(want_raw)

with tempfile.TemporaryDirectory() as td:
    path = os.path.join(td, "edit.xlsx")
    X.to_workbook(src).save(path)
    wb = openpyxl.load_workbook(path)
    members = X.LIST_BLOCKS["carbonation"][1]
    for i, fl in enumerate(want["floors"], 1):
        ws = wb[f"표본층{i}"]
        start = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 10).value == "carbonation") + 2
        end = start
        while ws.cell(end, 1).value != "(끝)":
            end += 1
        rows = [(m, rec) for m, cell in zip(members, fl["carbonation"]) for rec in cell]
        if len(rows) > end - start:                       # 줄이 모자라면 끼워 넣기 (사람이 행 삽입하듯)
            ws.insert_rows(end, len(rows) - (end - start))
        for r in range(start, start + max(len(rows), end - start) + 1):
            for c in range(1, X.EXACT_COL + 1):
                ws.cell(r, c).value = None
        for k, (m, rec) in enumerate(rows):               # 화면에 보이는 자리수(소수 셋째)로 입력
            ws.cell(start + k, 1, m)
            for j, x in enumerate(rec, 2):
                ws.cell(start + k, j, round(x, 3))
        ws.cell(start + len(rows), 1, "(끝)")
    wb.save(path)
    out = F.write(X.from_workbook(openpyxl.load_workbook(path, data_only=True)))

if out == want_raw:
    print(f"OK  엑셀에서 탄산화만 고쳐 만든 .sfm == SafetyMan 수정본 ({len(out)}바이트)")
    sys.exit(0)
diff = next((i for i in range(min(len(out), len(want_raw))) if out[i] != want_raw[i]), None)
print(f"DIFF 길이 {len(out)} vs {len(want_raw)}, 첫 차이 {diff}")
sys.exit(1)
